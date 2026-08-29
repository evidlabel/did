"""Tests for file_utils."""

import pytest
from docx import Document

from did.core.anonymizer import Anonymizer
from did.utils.file_utils import anonymize_file, extract_text


@pytest.fixture
def temp_files(tmp_path):
    md_file = tmp_path / "test.md"
    md_file.write_text("# Heading\n**Bold** text")
    txt_file = tmp_path / "test.txt"
    txt_file.write_text("Plain text")
    tex_file = tmp_path / "test.tex"
    tex_file.write_text(
        "\\documentclass{article} \\begin{document} Hello \\end{document}"
    )
    bib_file = tmp_path / "test.bib"
    bib_file.write_text(
        "@article{test, title={Test Title by John Doe}, author={John Doe}}"
    )
    return md_file, txt_file, tex_file, bib_file


def test_extract_text_md(temp_files):
    md_file, _, _, _ = temp_files
    text = extract_text(md_file)
    assert "# Heading" in text
    assert "**Bold** text" in text


def test_extract_text_txt(temp_files):
    _, txt_file, _, _ = temp_files
    text = extract_text(txt_file)
    assert text == "Plain text"


def test_extract_text_tex(temp_files):
    _, _, tex_file, _ = temp_files
    text = extract_text(tex_file)
    assert "Hello" in text
    assert "\\documentclass" not in text


def test_extract_text_bib(temp_files):
    _, _, _, bib_file = temp_files
    text = extract_text(bib_file)
    assert "Test Title by John Doe" in text
    assert "John Doe" in text


def test_extract_text_docx_preserves_paragraph_and_table_order(tmp_path):
    path = tmp_path / "case.docx"
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "John Doe"
    table.cell(0, 1).text = "Counsel"
    document.add_paragraph("After table")
    document.save(path)

    assert extract_text(path) == ("Before table\nJohn Doe\tCounsel\nAfter table")


def test_extract_text_unsupported(tmp_path):
    unsupported = tmp_path / "test.odt"
    unsupported.write_text("Dummy")
    with pytest.raises(ValueError, match="Unsupported file type: .odt"):
        extract_text(unsupported)


def test_anonymize_file_md(temp_files):
    md_file, _, _, _ = temp_files
    anonymizer = Anonymizer(language="en")
    anonymizer.detect_entities([extract_text(md_file)])
    output = md_file.with_stem("output")
    counts = anonymize_file(md_file, anonymizer, output)
    assert output.exists()
    assert counts["person_found"] == 0  # No persons in sample
    assert counts["person_replaced"] == 0


def test_anonymize_file_txt(temp_files):
    _, txt_file, _, _ = temp_files
    anonymizer = Anonymizer(language="en")
    anonymizer.detect_entities([extract_text(txt_file)])
    output = txt_file.with_stem("output")
    counts = anonymize_file(txt_file, anonymizer, output)
    assert output.exists()
    assert counts["person_found"] == 0
    assert counts["person_replaced"] == 0


def test_anonymize_file_tex(temp_files):
    _, _, tex_file, _ = temp_files
    anonymizer = Anonymizer(language="en")
    anonymizer.detect_entities([extract_text(tex_file)])
    output = tex_file.with_stem("output")
    counts = anonymize_file(tex_file, anonymizer, output)
    assert output.exists()
    assert counts["person_found"] == 0
    assert counts["person_replaced"] == 0


def test_anonymize_file_bib(temp_files):
    _, _, _, bib_file = temp_files
    anonymizer = Anonymizer(language="en")
    anonymizer.detect_entities([extract_text(bib_file)])
    output = bib_file.with_stem("output")
    counts = anonymize_file(bib_file, anonymizer, output)
    assert output.exists()
    with open(output) as f:
        bib_content = f.read()
        assert "#(P1V" in bib_content
    assert counts["person_replaced"] == 2


def test_anonymize_file_unsupported(tmp_path):
    unsupported = tmp_path / "test.odt"
    unsupported.write_text("Dummy")
    anonymizer = Anonymizer(language="en")
    output = unsupported.with_stem("output")
    with pytest.raises(ValueError, match="Unsupported file type: .odt"):
        anonymize_file(unsupported, anonymizer, output)


def test_organization_reaches_the_vars_file(tmp_path):
    """Regression: #(O1V1) in the body with no #let O1V1 to resolve it.

    ``export_to_typst`` kept its own category list, which omitted
    ``organization`` while the replacer's included it. The exported document
    referenced a variable nothing defined — the Typst compile failed and the
    organization's real value never reached the key set at all.
    """
    from did.utils.file_utils import export_to_typst

    source = tmp_path / "org.md"
    source.write_text("The contract with Acme Holdings was signed today.")
    anonymizer = Anonymizer(language="en")
    anonymizer.load_replacements(
        {"ORGANIZATION": [{"id": "ORGANIZATION_1", "variants": ["Acme Holdings"]}]}
    )

    main_path = tmp_path / "org.typ"
    export_to_typst(source, anonymizer, main_path)

    body = main_path.read_text()
    assert "#(O1V1)" in body
    assert "Acme Holdings" not in body
    # Every token in the body must be defined in the vars file.
    assert '#let O1V1 = "Acme Holdings"' in (tmp_path / "org_vars.typ").read_text()
    assert "#let O1V1" in (tmp_path / "org_fakevars.typ").read_text()
