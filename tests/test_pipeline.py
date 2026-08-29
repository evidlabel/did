"""Qt-free tests for the gdid pipeline (no PySide6, no spaCy)."""

import zipfile

import pytest
from docx import Document

from gdid import pipeline


class FakeAnonymizer:
    """Stand-in for did's Anonymizer that needs no spaCy model.

    Records calls and pseudonymizes by a trivial substring replacement so the
    pipeline's plumbing can be tested fast and deterministically.
    """

    def __init__(self, language="en"):
        self.language = language
        self.detected = None
        self.loaded = None

    def detect_entities(self, texts):
        self.detected = list(texts)

    def generate_yaml(self):
        return 'PERSON:\n  - id: "P1"\n    variants:\n      - "John Doe"\n'

    def load_replacements(self, config):
        self.loaded = config

    def anonymize(self, text):
        return text.replace("John Doe", "#(P1V1)"), {}


# ------------------------------------------------------------- collect_inputs ---
def test_collect_inputs_filters_supported(tmp_path):
    (tmp_path / "a.md").write_text("x")
    (tmp_path / "b.txt").write_text("x")
    (tmp_path / "c.pdf").write_bytes(b"%PDF-1.4")
    (tmp_path / "d.docx").write_bytes(b"docx collection does not parse files")
    files, temp_dirs = pipeline.collect_inputs([str(tmp_path)])
    names = sorted(f.name for f in files)
    assert names == ["a.md", "b.txt", "c.pdf", "d.docx"]
    assert temp_dirs == []


def test_collect_inputs_extracts_zip(tmp_path):
    src = tmp_path / "doc.md"
    src.write_text("hello")
    zip_path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(zip_path, "w") as zf:
        zf.write(src, arcname="doc.md")
    files, temp_dirs = pipeline.collect_inputs([str(zip_path)])
    assert [f.name for f in files] == ["doc.md"]
    assert len(temp_dirs) == 1
    assert temp_dirs[0].is_dir()


def test_collect_inputs_single_files(tmp_path):
    md = tmp_path / "a.md"
    md.write_text("x")
    files, temp_dirs = pipeline.collect_inputs([str(md)])
    assert files == [md]
    assert temp_dirs == []


# ----------------------------------------------------------------- extract_one ---
def test_extract_one_md(tmp_path):
    f = tmp_path / "a.md"
    f.write_text("# Heading\nbody text")
    text = pipeline.extract_one(f)
    assert "Heading" in text
    assert "body text" in text


def test_extract_one_txt(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("plain content")
    assert pipeline.extract_one(f) == "plain content"


def test_extract_one_docx(tmp_path):
    path = tmp_path / "letter.docx"
    document = Document()
    document.add_paragraph("Dear John Doe")
    document.save(path)

    assert pipeline.extract_one(path) == "Dear John Doe"


# --------------------------------------------------------------- detect_to_yaml ---
def test_detect_to_yaml_uses_injected_factory():
    texts = ["John Doe was here.", "Also John Doe."]
    anonymizer, yaml_str = pipeline.detect_to_yaml(
        texts, "en", anonymizer_factory=FakeAnonymizer
    )
    assert isinstance(anonymizer, FakeAnonymizer)
    assert anonymizer.language == "en"
    assert anonymizer.detected == texts
    assert "John Doe" in yaml_str


# --------------------------------------------------------------- pseudonymize_all ---
def test_pseudonymize_all_maps_each_text():
    anonymizer = FakeAnonymizer()
    texts = {"doc1": "John Doe met X.", "doc2": "Hi John Doe."}
    result = pipeline.pseudonymize_all(anonymizer, "PERSON: []", texts)
    assert result == {"doc1": "#(P1V1) met X.", "doc2": "Hi #(P1V1)."}
    assert anonymizer.loaded is not None  # load_replacements was called


def test_pseudonymize_all_rejects_bad_yaml():
    with pytest.raises(ValueError, match="Invalid YAML"):
        pipeline.pseudonymize_all(FakeAnonymizer(), "key: [unclosed", {"d": "x"})


def test_pseudonymize_all_rejects_empty_yaml():
    with pytest.raises(ValueError, match="empty"):
        pipeline.pseudonymize_all(FakeAnonymizer(), "", {"d": "x"})


# ------------------------------------------------------------------ verify_outputs ---
CONFIG_YAML = 'PERSON:\n  - id: "P1"\n    variants:\n      - "John Doe"\n'


def test_verify_outputs_is_clean_for_fully_pseudonymized_text():
    report = pipeline.verify_outputs(CONFIG_YAML, {"d": "#(P1V1) met X."})
    assert report.is_clean
    assert report.deep is False


def test_verify_outputs_reports_a_surviving_surname():
    """``Doesagen`` keeps the surname readable; the shallow tiers catch it."""
    report = pipeline.verify_outputs(CONFIG_YAML, {"d": "#(P1V1) met Doesagen."})
    assert [f.text for f in report.leaks] == ["Doe"]
    assert not report.is_clean


def test_verify_outputs_separates_deliberate_exclusions():
    report = pipeline.verify_outputs(
        "PERSON: []", {"d": "Signed by Jane Roe."}, ["Jane Roe"]
    )
    assert report.leaks == []
    assert [f.text for f in report.retained] == ["Jane Roe"]


def test_verify_outputs_tolerates_an_empty_config():
    report = pipeline.verify_outputs("", {"d": "#(P1V1) only."})
    assert report.is_clean


def test_verify_outputs_degrades_when_the_anonymizer_cannot_rescan():
    """A stand-in cannot re-run detection; the report must say so, not claim it."""
    report = pipeline.verify_outputs(
        CONFIG_YAML, {"d": "#(P1V1)."}, anonymizer=FakeAnonymizer()
    )
    assert report.deep is False
    assert "does not support re-detection" in report.deep_error


def test_exclude_not_names_removes_matching_person_identity():
    yaml_text = (
        "PERSON:\n"
        '  - id: "PERSON_1"\n'
        '    variants: ["Virksomhedstype", "Virksomhedstypen"]\n'
        '  - id: "PERSON_2"\n'
        '    variants: ["Jane Doe"]\n'
        "LOCATION: []\n"
    )

    filtered = pipeline.exclude_not_names(yaml_text, ["virksomhedstype"])
    data = pipeline.parse_yaml(filtered)

    assert [entity["id"] for entity in data["PERSON"]] == ["PERSON_2"]
    assert data["PERSON"][0]["variants"] == ["Jane Doe"]


def test_resolve_selection_placeholders_restores_variant_inside_larger_phrase():
    yaml_text = 'PERSON:\n  - id: "PERSON_1"\n    variants: ["John", "J."]\n'

    assert (
        pipeline.resolve_selection_placeholders("#(P1V1) Doe", yaml_text) == "John Doe"
    )
    assert pipeline.resolve_selection_placeholders("#(P1V2) Doe", yaml_text) == "J. Doe"
    assert (
        pipeline.resolve_selection_placeholders("[PERSON 1] Doe", yaml_text)
        == "John Doe"
    )
    assert (
        pipeline.resolve_selection_placeholders("#(P99V1) Doe", yaml_text)
        == "#(P99V1) Doe"
    )


def test_merge_person_entities_moves_whole_identity():
    yaml_text = (
        "PERSON:\n"
        '  - id: "PERSON_1"\n'
        '    variants: ["J. Doe"]\n'
        '  - id: "PERSON_2"\n'
        '    variants: ["John Doe"]\n'
    )

    merged = pipeline.merge_person_entities(yaml_text, "PERSON_1", "PERSON_2")
    people = pipeline.parse_yaml(merged)["PERSON"]

    assert len(people) == 1
    assert people[0]["id"] == "PERSON_2"
    assert people[0]["variants"] == ["John Doe", "J. Doe"]


def test_merge_person_entities_moves_one_variant():
    yaml_text = (
        "PERSON:\n"
        '  - id: "PERSON_1"\n'
        '    variants: ["J. Doe", "Jane Doe"]\n'
        '  - id: "PERSON_2"\n'
        '    variants: ["John Doe"]\n'
    )

    merged = pipeline.merge_person_entities(yaml_text, "PERSON_1", "PERSON_2", "J. Doe")
    people = pipeline.parse_yaml(merged)["PERSON"]

    assert people[0]["variants"] == ["Jane Doe"]
    assert people[1]["variants"] == ["John Doe", "J. Doe"]


# ----------------------------------------------------------------- save_outputs ---
def _stub_export(monkeypatch):
    """Replace export_to_typst with a recorder that writes a token file."""
    calls = []

    def fake_export(input_path, anonymizer, main_path, **kwargs):
        calls.append((input_path, main_path, kwargs))
        main_path.parent.mkdir(parents=True, exist_ok=True)
        main_path.write_text("#(P1V1) body\n", encoding="utf-8")

    monkeypatch.setattr(pipeline, "export_to_typst", fake_export)
    return calls


def test_save_outputs_multi(tmp_path, monkeypatch):
    calls = _stub_export(monkeypatch)
    files = [tmp_path / "a.md", tmp_path / "b.md"]
    for f in files:
        f.write_text("x")
    sub = pipeline.save_outputs(
        files, FakeAnonymizer(), "PERSON: []", "multi", tmp_path
    )
    assert sub == tmp_path / "pseudonymized"
    assert (sub / "a_pseudonymized.typ").exists()
    assert (sub / "b_pseudonymized.typ").exists()
    assert (sub / "config.yaml").read_text() == "PERSON: []"
    # shared vars filenames passed through to every export call
    for _, _, kwargs in calls:
        assert kwargs["vars_filename"].endswith("shared_vars.typ")
        assert kwargs["fakevars_filename"].endswith("shared_fakevars.typ")


def test_save_outputs_single_combines(tmp_path, monkeypatch):
    calls = _stub_export(monkeypatch)
    files = [tmp_path / "a.md", tmp_path / "b.md"]
    for f in files:
        f.write_text("x")
    sub = pipeline.save_outputs(
        files, FakeAnonymizer(), "PERSON: []", "single", tmp_path
    )
    assert sub == tmp_path / "single_pseudonymized"
    combined = (sub / "combined.typ").read_text()
    assert combined.startswith('#import "shared_vars.typ": *')
    assert "= a.md" in combined
    assert "= b.md" in combined
    # single-combine must request token-only bodies (no import header to strip)
    assert all(kwargs.get("write_imports") is False for _, _, kwargs in calls)


def test_save_outputs_bad_mode(tmp_path):
    with pytest.raises(ValueError, match="Unknown save mode"):
        pipeline.save_outputs([], FakeAnonymizer(), "x", "weird", tmp_path)


# ------------------------------------------------------------ to_written_out ---
def test_to_written_out_all_prefixes():
    text = (
        "#(P1V1) #(E2V1) #(A1V1) #(PH3V1) #(DT4V2) #(ID5V1) #(CD6V1) #(GN7V9) #(URL8V1)"
    )
    assert pipeline.to_written_out(text) == (
        "[PERSON 1] [EMAIL 2] [ADDRESS 1] [PHONE 3] [DATE 4] "
        "[ID 5] [CODE 6] [NUMBER 7] [URL 8]"
    )


def test_to_written_out_collapses_variants():
    assert (
        pipeline.to_written_out("#(P1V1) met #(P1V2).") == "[PERSON 1] met [PERSON 1]."
    )


def test_to_written_out_leaves_other_text_alone():
    text = "plain [brackets] and #(X1V1) stay as-is"
    assert pipeline.to_written_out(text) == text


def test_to_redacted_replaces_all_supported_placeholders():
    text = "#(P1V1) called #(PH2V1) on #(DT3V1); #(X1V1) stays."
    assert pipeline.to_redacted(text) == (
        "[REDACTED] called [REDACTED] on [REDACTED]; #(X1V1) stays."
    )


def test_to_synthetic_is_stable_and_groups_variants():
    yaml_text = (
        "PERSON:\n"
        '  - id: "PERSON_1"\n'
        '    variants: ["John Doe", "John"]\n'
        '  - id: "PERSON_2"\n'
        '    variants: ["Jane Doe"]\n'
    )
    text = "#(P1V1) met #(P1V2) and #(P2V1); unknown #(P9V1)."

    first = pipeline.to_synthetic(text, yaml_text, "en", "case-123")
    second = pipeline.to_synthetic(text, yaml_text, "en", "case-123")

    assert first == second
    values = first.removesuffix(".").split(" ")
    assert "#(P1V1)" not in first
    assert "#(P1V2)" not in first
    assert "#(P2V1)" not in first
    assert "#(P9V1)" in first
    # Both variants of identity 1 render as precisely the same generated name.
    left, remainder = first.split(" met ", 1)
    repeated, _ = remainder.split(" and ", 1)
    assert left == repeated
    assert values
